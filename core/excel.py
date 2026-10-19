"""excel.py - the income summary as an Excel workbook (same layout as the screen / office template)."""
from __future__ import annotations

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

NUMFMT = '#,##0.00;-#,##0.00;""'


def income_xlsx(inc: dict, client: str = "", ded: dict | None = None, docs: list | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Income Summary"
    F = "Arial"
    bold, red, head = Font(name=F, size=10, bold=True), Font(name=F, size=10, color="C00000"), Font(name=F, size=10, bold=True, color="FFFFFF")
    norm, grey_t = Font(name=F, size=10), Font(name=F, size=10, color="808080", italic=True)
    tot_fill, grand_fill = PatternFill("solid", fgColor="EDEDED"), PatternFill("solid", fgColor="D9D9D9")
    head_fill, edit_fill = PatternFill("solid", fgColor="003366"), PatternFill("solid", fgColor="FFF2CC")
    thin = Side(style="thin", color="9AA3B2")
    ws["A1"] = "Income Summary" + (f" — {client}" if client else "")
    ws["A1"].font = Font(name=F, size=12, bold=True, color="003366")
    cy = inc.get("cy_year") or "CY"
    hdr = ["Document/s", str(inc.get("py_year") or "PY"), f"{cy} per document", f"{cy} per preparer", "Fed. W/H", "Comments"]
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(row=3, column=c, value=h)
        cell.font, cell.fill = head, head_fill
        cell.alignment = Alignment(horizontal="right" if 1 < c < 6 else "left", wrap_text=True)
    r = 4
    blk_fill, blk_tot = PatternFill("solid", fgColor="FFFFFF"), PatternFill("solid", fgColor="E3F1E8")
    bracket = Side(style="thick", color="294F4C")

    def put(label, py=None, dcy=None, cy=None, fit=None, note="", font=norm, fill=None, differs=False, fit_changed=False, block=False, total=False):
        nonlocal r
        vals = [label, py, dcy, cy, fit, note]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.font = font if c != 6 else (blue if font is blue else Font(name=F, size=9, color="555555"))
            if c == 3 and font is not bold:
                cell.font = norm if font is not grey_t else grey_t
            if 1 < c < 6:
                cell.number_format = NUMFMT
            if fill:
                cell.fill = fill
            elif block:
                cell.fill = blk_fill
            if (differs and c == 4) or (fit_changed and c == 5):
                cell.fill = edit_fill
            cell.border = Border(bottom=Side(style="medium", color="294F4C") if total else thin)    # a thicker line after each total
            cell.alignment = Alignment(vertical="top", wrap_text=(c in (1, 6)))
        r += 1

    green = Font(name=F, size=10, color="1F5FBF")                 # typed by the preparer: blue, not bold
    blue = Font(name=F, size=10, color="1F5FBF")

    def sections(data):
      for s in data["sections"]:
        for d in s["rows"]:
            f = grey_t if (d.get("excluded") or d.get("memo")) else (green if (d.get("added") or d.get("edited")) else norm)
            fit_changed = d.get("edited") and abs((d.get("read_fit") or 0) - (d.get("fit") or 0)) > 0.005
            note = ("From the document inventory. " if d.get("inv") else "Added by the preparer. " if d.get("added") else "") + (d.get("note") or "")
            if fit_changed:
                note = f"W/H changed from {(d.get('read_fit') or 0):,.2f}. " + note
            if d.get("note_edited"):
                f = green if f is norm else f
            put(d["doc"], None, None if d.get("excluded") else d.get("doc_cy"), None if d.get("excluded") else d.get("cy"),
                d.get("fit"), note, f, differs=abs((d.get("doc_cy") or 0) - (d.get("cy") or 0)) > 0.005 and d.get("doc_cy") is not None,
                fit_changed=fit_changed, block=True)
        for x in s["red"]:
            if x["amount"] is None and not s["rows"]:
                continue
            f = blue                                            # add-backs / minus losses carried over: blue
            put(x["label"], None, None, x["amount"], None, x["note"], f, block=True)
        put(s["label"], s.get("py"), s.get("doc_cy"), s.get("cy"), s.get("fit"), s.get("note", ""), bold, None, total=True)

    def r_gap():
        nonlocal r
        r += 1

    sections(inc)
    put("Total Income", inc.get("py_total"), inc.get("doc_total"), inc.get("cy_total"), inc.get("cy_fit"),
        inc.get("total_note", ""), bold, grand_fill)
    put("Federal income tax withheld", inc.get("py_fit"), None, None, inc.get("cy_fit"),
        "PY: Form 1040 line 25d · CY: W-2 box 2, 1099 box 4, SSA-1099 box 6", bold, tot_fill)
    if ded and ded.get("sections"):
        r += 1
        ws.cell(row=r, column=1, value="Deduction Summary").font = Font(name=F, size=12, bold=True, color="003366")
        r += 1
        sections(ded)
        put("Total Deductions", ded.get("py_total"), ded.get("doc_total"), ded.get("cy_total"), None, "", bold, grand_fill)
    for col, w in zip("ABCDEF", (52, 15, 16, 16, 15, 70)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A4"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToHeight = 0
    if docs:
        documents_sheet(wb, docs)
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def documents_sheet(wb, docs: list, title: str = "Documents"):
    """The document inventory: verified tick, file, issuer, year, read by the program, amount / line, note, AI."""
    ws = wb.create_sheet(title)
    F = "Arial"
    head, head_fill = Font(name=F, size=10, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="003366")
    ok_fill, thin = PatternFill("solid", fgColor="E6F4EA"), Side(style="thin", color="9AA3B2")
    hdr = ["Verified", "File name", "Issuer", "Year", "Program read", "Amount", "Added to", "Note", "AI"]
    for c, h in enumerate(hdr, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill = head, head_fill
    for i, d in enumerate(docs, 2):
        vals = ["✓" if d.get("verified") else "", d.get("file"), d.get("issuer"), d.get("year"), d.get("read_lbl"),
                d.get("cy"), d.get("sec_lbl") or d.get("feeds"), d.get("note"), "✓" if d.get("ai") else ""]
        for c, v in enumerate(vals, 1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.font = Font(name=F, size=10)
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=c in (2, 7, 8))
            if c == 6:
                cell.number_format = NUMFMT
            if d.get("verified"):
                cell.fill = ok_fill
    for col, w in zip("ABCDEFGHI", (9, 46, 22, 7, 13, 14, 30, 60, 5)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"
    return ws
