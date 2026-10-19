"""verify.py - screen 2: check the finished 2025 return against what the program read (all by program, no AI).

  1. Category ties   - every income summary total (2025 column, with your edits / added lines) vs the return line
  2. Per document    - each document amount found on the return (Federal Statements, 8949, Schedule B / D, 6781 ...)
  3. Schedule D      - by Form 8949 box: the documents' rows vs the return's Schedule D lines (1b/2/3, 4, 8b/9/10, 11, 13)
  4. 8949 boxes      - box checked on each Form 8949 page vs the box each 1099 / Koinly row belongs to
  5. Last year       - forms filed, payers / K-1 entities and filing status: 2024 return vs 2025 return

Checklist (master list + this return's lines) and the Excel "Save review" are here too.
"""
from __future__ import annotations

import json
import re
from datetime import datetime
from io import BytesIO
from pathlib import Path

from .pdfutil import open_pdf, page_text
from .py_return import read_prior_return, read_py_income, read_schedule_e, FORM_SHORT

HERE = Path(__file__).resolve().parent.parent
MASTER = HERE / "checklist_master.json"
DEFAULT_CHECKLIST = [
    "All FYIs / diagnostics cleared in UT",
    "Filing status makes sense — also last year's (MFS vs MFJ?)",
    "Penalties (FTF / FTP / 2210): calculating? should they be?",
    "Federal and state extension and estimated payments entered",
    "Next-year state estimates generated if needed (e.g. CA Q4)",
    "Health insurance: 1095-A received → 8962 (allocated if shared)",
    "Entities: initial / final / dissolved boxes; CA LLC $800; dissolution follow-up noted",
    "MFS: spouse's deduction method (standard / itemized) matches; 8958 if community property state",
    "Check carryovers",
]

AMT = re.compile(r"^\(?-?\$?[\d,]*\d(?:\.\d{1,2})?\)?$")


def _n(s):
    s = s.replace("$", "").replace(",", "").strip()
    neg = s.startswith("(") or s.startswith("-")
    try:
        v = float(s.strip("()-"))
    except ValueError:
        return None
    return -v if neg else v


def _rows(words, tol=3.5):
    rows = []
    for w in sorted(words, key=lambda w: (w[1], w[0])):
        if rows and abs(rows[-1][0] - w[1]) <= tol:
            rows[-1][1].append(w)
        else:
            rows.append([w[1], [w]])
    return [sorted(r[1], key=lambda w: w[0]) for r in rows]


# ---------------------------------------------------------------- master checklist (survives restarts)
def master_load() -> list[str]:
    try:
        v = json.loads(MASTER.read_text(encoding="utf-8"))
        return [str(x) for x in v.get("items", [])]
    except (OSError, ValueError):
        return list(DEFAULT_CHECKLIST)


def master_save(items: list[str]) -> None:
    MASTER.write_text(json.dumps(dict(items=items, updated=datetime.now().strftime("%Y-%m-%d %H:%M")), indent=2,
                                 ensure_ascii=False), encoding="utf-8")


# ---------------------------------------------------------------- reading the 2025 return
def _sched_d(doc, texts) -> dict:
    """Schedule D column (h) by line: {'1b': .., '2': .., '3': .., '4': .., '7': .., '8b': .., '15': .., '16': ..}."""
    out = {}
    for i, t in enumerate(texts):
        if not re.search(r"SCHEDULE D\s*\(Form 1040\)|Capital Gains and Losses", t) or "Short-Term" not in t:
            continue
        page = doc[i]
        W = page.rect.width
        words = page.get_text("words")
        ids = [w for w in words if w[0] < 0.12 * W and re.fullmatch(r"\d{1,2}[ab]?", w[4])]
        for ws in _rows(words):
            amts = [w for w in ws if w[2] > 0.86 * W and AMT.match(w[4]) and any(c.isdigit() for c in w[4])]
            if not amts:
                continue
            y = ws[0][1]
            near = [w for w in ids if -3 <= y - w[1] <= 22]
            if not near:
                continue
            line = max(near, key=lambda w: w[1])[4]
            if line not in out:
                v = _n(amts[-1][4])
                if v is not None:
                    out[line] = v
        if out:
            # page 2: line 16
            for j in range(i + 1, min(i + 3, len(texts))):
                m = re.search(r"\b16\s+Combine lines 7 and 15.*?16\s+(\(?-?[\d,]+\)?)", re.sub(r"\s+", " ", texts[j]))
                if m:
                    out["16"] = _n(m.group(1))
                    break
            break
    return out


def _f8949_boxes(doc, texts) -> list[dict]:
    """Box checked on each Form 8949 page (A-L) and the page total (h)."""
    res = []
    for i, t in enumerate(texts):
        if "8949" not in t or not re.search(r"Box [A-L]\b", t) or not re.search(r"Part I", t):
            continue
        page = doc[i]
        words = page.get_text("words")
        boxes = [(w[1], w[3], w[4][1]) for w in words if re.fullmatch(r"\([A-L]\)", w[4])]
        if not boxes:
            continue
        marks = [w for w in words if w[4] in ("X", "x", "✓", "✔", "■") and w[0] < 0.2 * page.rect.width]
        checked = None
        for m in marks:
            cy = (m[1] + m[3]) / 2
            hit = [b for b in boxes if b[0] - 3 <= cy <= b[1] + 3]
            if hit:
                checked = hit[0][2]
                break
        tot = None
        for ws in _rows(words):
            txt = " ".join(w[4] for w in ws)
            if re.search(r"checked\)|Totals\.|Add the amounts", txt):
                amts = [w for w in ws if AMT.match(w[4]) and any(c.isdigit() for c in w[4]) and w[0] > 0.45 * page.rect.width]
                if amts:
                    tot = _n(amts[-1][4])
        res.append(dict(page=i + 1, box=checked, total=tot))
    return res


def read_cy_return(path) -> dict:
    py = read_prior_return(path)
    inc = read_py_income(path, py, [])
    doc, _ = open_pdf(path)
    with doc:
        texts = [page_text(doc, i) for i in range(len(doc))]
        schd = _sched_d(doc, texts)
        f8949 = _f8949_boxes(doc, texts)
    flat = [re.sub(r"\s+", " ", t) for t in texts]
    try:
        sche = read_schedule_e(path)
    except Exception:  # noqa: BLE001
        sche = []
    return dict(sche=sche, file=Path(path).name, year=py.get("year"), filing_status=py.get("filing_status"), forms=py.get("forms", []),
                expected=py.get("expected", []), inc=inc, schd=schd, f8949=f8949, texts=flat, pages=len(texts),
                summary=py.get("summary"))


# ---------------------------------------------------------------- 1095-A / 1098-T (not income: checked on their own)
MEMO_FORMS = {"1095a": ("1095-A", ["8962", "Premium Tax Credit (PTC)", "Annual and Monthly Contribution"], "Form 8962 line 25"),
              "1098t": ("1098-T", ["8863", "Education Credits", "Refundable American opportunity"], "Form 8863"),
              "1099sa": ("1099-SA", ["8889", "Health Savings Accounts", "HSA"], "Form 8889 line 14a")}


def memo_items(docs: list) -> list:
    """What the 1095-A / 1098-T documents say, to check on screen 2: {form, doc, what, amount, kind}."""
    out = []
    for d in docs:
        p = d.get("parsed")
        if not (isinstance(p, dict) and p.get("memo_only")):
            continue
        det = p.get("detail") or {}
        if d["family"] == "1095a" and det.get("aptc") is not None:
            out.append(dict(kind="1095a", form="1095-A", doc=d["file"], what="Advance PTC (line 33C)", amount=det["aptc"]))
        if d["family"] == "1099sa" and det.get("gross"):
            out.append(dict(kind="1099sa", form="1099-SA", doc=d["file"], what="Box 1 gross distribution", amount=det["gross"]))
        if d["family"] == "1098t":
            for bx, what in (("box1", "Box 1 payments"), ("box5", "Box 5 scholarships")):
                if det.get(bx):
                    out.append(dict(kind="1098t", form="1098-T", doc=d["file"], what=what, amount=det[bx]))
    return out


def memo_return(cy: dict, docs: list) -> list:
    """1095-A / 1098-T vs the 2025 return: the form is there and the amount is printed on it."""
    out = []
    for it in memo_items(docs):
        form, marks, line = MEMO_FORMS[it["kind"]]
        pages = [i + 1 for i, t in enumerate(cy.get("texts") or []) if marks[0] in t and sum(m in t for m in marks[1:]) >= 1
                 and re.search(r"Form\s+" + marks[0], t)]
        x = dict(it, where=line)
        if not pages:
            x.update(status="no form", note=f"No {line.split(' line')[0]} on the return — {form} received")
        else:
            amt = f"{round(it['amount']):,}"
            hit = [p for p in pages if re.search(r"(?<![\d,])" + re.escape(amt) + r"(?![\d,])", cy["texts"][p - 1])]
            x.update(status="found" if hit else "not found", note=f"pages {', '.join(map(str, (hit or pages)[:6]))}")
        out.append(x)
    return out


# ---------------------------------------------------------------- checks
CAT_LINE = {"w2": "1z wages", "interest": "2b taxable interest", "dividends": "3b ordinary dividends", "ira": "4b IRA",
            "pension": "5b pensions", "ss": "6b social security", "st": "Sch D line 7", "lt": "Sch D line 15",
            "schc": "Sch 1 line 3 (Sch C)", "other": "Sch 1 other income", "rental": "Sch E line 26", "k1": "Sch E line 32"}
SCHD_OF = {"A": "1b", "G": "1b", "B": "2", "H": "2", "C": "3", "I": "3", "D": "8b", "J": "8b", "E": "9", "K": "9", "F": "10", "L": "10"}
SCHD_LBL = {"1b": "1b — Box A / G", "2": "2 — Box B / H", "3": "3 — Box C / I", "4": "4 — 6781 / 4684 / 6252 (ST)", "5": "5 — K-1 ST",
            "8b": "8b — Box D / J", "9": "9 — Box E / K", "10": "10 — Box F / L", "11": "11 — 6781 / 4797 (LT)", "12": "12 — K-1 LT",
            "13": "13 — capital gain distributions"}


def _status(doc_v, ret_v, tol):
    if doc_v is None and (ret_v is None or abs(ret_v) < 0.5):
        return None
    if doc_v is None:
        return "return only"
    if ret_v is None:
        return "missing" if abs(doc_v) > tol else "ok"
    return "ok" if abs(doc_v - ret_v) <= tol else "diff"


def _schd_line(label: str, term: str) -> str | None:
    m = re.search(r"\(([A-L])(?:/([A-L]))?\)", label)
    if m:
        return SCHD_OF.get(m.group(1))
    if "6781" in label:
        return "4" if term == "st" else "11"
    if "Sch D 13" in label:
        return "13"
    return None


def _akey(n) -> str:
    n = re.sub(r"[^a-z0-9 ]", " ", str(n or "").lower())
    n = re.sub(r"\b(avenue|ave|street|st|drive|dr|road|rd|place|pl|north|n|south|s|east|e|west|w)\b", "", n)
    return re.sub(r"\s+", " ", n).strip()


def sche_return(cy: dict, sche: list | None) -> list:
    """Schedule E box (each property, each line) vs Schedule E page 1 of the 2025 return."""
    from core.py_return import read_schedule_e, SCHE_LINES
    ret = list(cy.get("sche") or [])
    out = []
    for p in sche or []:
        if p.get("placeholder") or not p.get("calculated"):
            continue
        cyv = {k: float(str(v).replace(",", "")) for k, v in (p.get("cy") or {}).items() if str(v).strip() not in ("",) and k in SCHE_LINES}
        tot = p.get("totals") or {}
        for k in ("20", "21"):
            if tot.get(k) is not None:
                cyv[k] = tot[k]
        if "22" in cyv:
            cyv["22"] = -abs(cyv["22"])
        pk = _akey(p.get("name"))
        hit = next((r for r in ret if pk and _akey(r.get("name")) and (_akey(r.get("name")) == pk or _akey(r.get("name")).startswith(pk + " ")
                                                                    or pk.startswith(_akey(r.get("name")) + " "))), None) or \
            next((r for r in ret if cyv.get("3") and abs((r["lines"].get("3") or 0) - cyv["3"]) <= 1), None)
        if hit:
            ret.remove(hit)
        for k in SCHE_LINES:
            a, b = cyv.get(k), (hit or {}).get("lines", {}).get(k)
            if a in (None, 0) and b in (None, 0):
                continue
            st = "missing" if hit is None else ("ok" if a is not None and b is not None and abs(a - b) <= 1.01 else "diff")
            out.append(dict(prop=p.get("name") or "", ret_prop=(hit or {}).get("name", ""), line=k, label=SCHE_LINES[k],
                            box=a, ret=b, diff=None if a is None or b is None else round((a or 0) - (b or 0), 2), status=st))
    for r in ret:                                         # on the return, not in the box
        out.append(dict(prop="", ret_prop=r.get("name", ""), line="", label="property on the return, not in the Schedule E box",
                        box=None, ret=r["lines"].get("21"), diff=None, status="return only"))
    return out


def run_checks(cy: dict, inc: dict, ded: dict, docs: list, py: dict, sche: list | None = None) -> dict:
    out = dict(cats=[], perdoc=[], schd=[], boxes=[], forms=[], payers=[], status=[], memo=memo_return(cy, docs),
               sche=sche_return(cy, sche))
    rinc = cy["inc"]
    # 1. category ties
    for s in inc["sections"]:
        k = s["key"]
        if k not in CAT_LINE:
            continue
        dv, rv = s.get("cy"), rinc.get(k)
        n = max(1, len([r for r in s["rows"] if r.get("cy") is not None]))
        st = _status(dv, rv, 1.01 * n)          # ±$1 per line (people round either way)
        if st:
            out["cats"].append(dict(cat=s["label"], line=CAT_LINE[k], docs=dv, ret=rv,
                                    diff=None if dv is None or rv is None else round(dv - rv, 2), status=st))
    dv, rv = inc.get("cy_fit"), rinc.get("fit")
    fit_rows = [r for s in inc["sections"] for r in s["rows"] if r.get("fit")]
    st = _status(dv or None, rv, 1.01 * max(1, len(fit_rows)))
    if st:
        out["cats"].append(dict(cat="Federal income tax withheld", line="25d", docs=dv or None, ret=rv,
                                diff=None if not dv or rv is None else round(dv - rv, 2), status=st))
    for s in (ded or {}).get("sections", []):
        rv = rinc.get(s["key"])
        st = _status(s.get("cy"), rv, 1.01 * max(1, len(s["rows"])))
        if st:
            out["cats"].append(dict(cat=s["label"], line="Sch 1 line 21" if s["key"] == "student_loan" else "",
                                    docs=s.get("cy"), ret=rv, diff=None if s.get("cy") is None or rv is None else round(s["cy"] - rv, 2),
                                    status=st))
    # 2. per document: the amount (rounded) printed anywhere on the return
    allt = cy["texts"]
    for s in list(inc["sections"]) + list((ded or {}).get("sections", [])):
        for r in s["rows"]:
            v = r.get("cy")
            if v is None or r.get("excluded") or r.get("memo") or abs(v) < 10:
                continue
            pats = {f"{round(v):,}", f"{abs(round(v)):,}", f"{v:,.2f}"}
            pages = [i + 1 for i, t in enumerate(allt) if any(re.search(r"(?<![\d,.])" + re.escape(p) + r"(?![\d,]|\.\d)", t) for p in pats)]
            out["perdoc"].append(dict(cat=s["label"], doc=r["doc"] or "(added line)", amount=v, pages=pages[:8], found=bool(pages),
                                      source=r.get("source", "read by program"), change=r.get("change_note", "")))
        for x in s.get("red", []):
            v = x.get("amount")
            if v is None or abs(v) < 10:
                continue
            pats = {f"{round(v):,}", f"{abs(round(v)):,}", f"{v:,.2f}"}
            pages = [i + 1 for i, t in enumerate(allt) if any(re.search(r"(?<![\d,.])" + re.escape(p) + r"(?![\d,]|\.\d)", t) for p in pats)]
            out["perdoc"].append(dict(cat=s["label"], doc=x["label"], amount=v, pages=pages[:8], found=bool(pages),
                                      source="added by preparer", change=""))
    # 3. Schedule D by box
    mine = {}
    for s in inc["sections"]:
        if s["key"] not in ("st", "lt"):
            continue
        for r in s["rows"]:
            if r.get("cy") is None or r.get("excluded") or r.get("memo"):
                continue
            line = _schd_line(r["doc"] or "", s["key"]) or ("5" if s["key"] == "st" and (r["doc"] or "").startswith("K-1") else
                                                            "12" if (r["doc"] or "").startswith("K-1") else None)
            if line is None and r.get("added"):
                line = "+" + s["key"]                    # your added line without a box: compared in the term total
            mine.setdefault(line or "?", []).append(r)
        for x in s.get("red", []):
            if x.get("amount") is not None:
                mine.setdefault("+" + s["key"], []).append(dict(doc=x["label"], cy=x["amount"]))
    sd = cy["schd"]
    for line in ["1b", "2", "3", "4", "5", "8b", "9", "10", "11", "12", "13", "+st", "+lt", "?"]:
        rows = mine.get(line, [])
        rv = sd.get(line)
        if line == "1b" and "1a" in sd:
            rv = (rv or 0) + sd["1a"]
        if line == "8b" and "8a" in sd:
            rv = (rv or 0) + sd["8a"]
        dv = sum(r["cy"] for r in rows) if rows else None
        if dv is None and rv is None:
            continue
        if line.startswith("+"):
            out["schd"].append(dict(line=("Short-term" if line == "+st" else "Long-term") + " — your added lines / add-backs (no box)",
                                    docs=dv, ret=None, diff=None, status="check", rows=[r["doc"] or "(added line)" for r in rows]))
            continue
        out["schd"].append(dict(line=SCHD_LBL.get(line, "not matched to a Schedule D line"), docs=dv, ret=rv,
                                diff=None if dv is None or rv is None else round(dv - rv, 2),
                                status=_status(dv, rv, 1.01 * max(1, len(rows))) if line != "?" else "check",
                                rows=[r["doc"] for r in rows]))
    # 4. 8949 boxes: box letters on the return vs the documents
    ret_boxes = sorted({b["box"] for b in cy["f8949"] if b["box"]})
    doc_boxes = {}
    for line, rows in mine.items():
        for r in rows:
            m = re.search(r"\(([A-L])(?:/([A-L]))?\)", r["doc"])
            if m:
                doc_boxes.setdefault(m.group(0), []).append(r["doc"])
    for b in cy["f8949"]:
        out["boxes"].append(dict(page=b["page"], box=b["box"] or "?", total=b["total"]))
    for lab, rows in doc_boxes.items():
        letters = [x for x in lab.strip("()").split("/") if x]
        on = [x for x in letters if x in ret_boxes]
        if len(letters) == 1:
            st = "ok" if on else ("no 8949" if not ret_boxes else "diff")
        else:
            st = "ok" if len(on) == 1 else ("diff" if not on else "check")
        sdl = SCHD_OF.get(letters[0])
        sdv = sd.get(sdl) if sdl else None
        out["boxes"].append(dict(page="", box=lab, docs=rows, status=st,
                                 note=("on the return as Box " + on[0]) if len(on) == 1 and len(letters) > 1 else
                                 ("" if on else f"no Form 8949 page with Box {' or '.join(letters)}"
                                  + (f" — Schedule D line {sdl} shows {sdv:,.0f}" if sdv is not None else ""))))
    # 5. last year vs this year
    pyf = {f["form"]: f for f in py.get("forms", [])}
    cyf = {f["form"]: f for f in cy.get("forms", [])}
    for name in list(pyf) + [x for x in cyf if x not in pyf]:
        short = FORM_SHORT.get(name.replace("Form ", ""), "")
        st = "both" if name in pyf and name in cyf else ("missing" if name in pyf else "new")
        out["forms"].append(dict(form=name, short=short, status=st))
    pye = [(e["kind"], e["name"]) for e in py.get("expected", [])]
    cye = [(e["kind"], e["name"]) for e in cy.get("expected", [])]
    from .py_return import _key
    for kind, name in pye:
        k = _key(name)
        hit = any(k and (k & _key(n2)) and max(len(t) for t in (k & _key(n2))) >= 4 for _, n2 in cye)
        out["payers"].append(dict(kind=kind, name=name, status="both" if hit else "missing"))
    for kind, name in cye:
        k = _key(name)
        if not any(k and (k & _key(n2)) and max(len(t) for t in (k & _key(n2))) >= 4 for _, n2 in pye):
            out["payers"].append(dict(kind=kind, name=name, status="new"))
    fp, fc = py.get("filing_status"), cy.get("filing_status")
    out["status"].append(dict(item="Filing status", py=fp or "not read", cy=fc or "not read",
                              status="check" if not fp or not fc else ("ok" if fp == fc else "changed")))
    tot_ret = rinc.get("total_return")
    out["status"].append(dict(item="Total income (1040 line 9)", py=(py.get("summary") and None), cy=tot_ret,
                              docs=inc.get("cy_total"), status=_status(inc.get("cy_total"), tot_ret, 2.0) or ""))
    return out


# ---------------------------------------------------------------- Save review (Excel)
def review_xlsx(client: str, cy_file: str, checks: dict | None, checklist: list, docs: list | None = None, ut: dict | None = None,
                comments: dict | None = None, expected: list | None = None, marks: dict | None = None, inc: dict | None = None,
                ded: dict | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    F = "Arial"
    head, bold, norm = Font(name=F, size=10, bold=True, color="FFFFFF"), Font(name=F, size=10, bold=True), Font(name=F, size=10)
    hfill = PatternFill("solid", fgColor="003366")
    ok_f, bad_f = PatternFill("solid", fgColor="E3F1E8"), PatternFill("solid", fgColor="FBE4E4")
    NUM = '#,##0.00;(#,##0.00)'
    mk = marks or {}
    tick = lambda k: "✓" if mk.get(k) else ""

    def table(ws, r, title, cols, rows, widths=None):
        ws.cell(r, 1, title).font = Font(name=F, size=11, bold=True, color="003366")
        r += 1
        for c, h in enumerate(cols, 1):
            x = ws.cell(r, c, h)
            x.font, x.fill = head, hfill
        r += 1
        for row in rows:
            for c, v in enumerate(row, 1):
                x = ws.cell(r, c, v)
                x.font = norm
                if isinstance(v, float):
                    x.number_format = NUM
                x.alignment = Alignment(vertical="top", wrap_text=isinstance(v, str) and len(v) > 40)
            st = str(row[-1]).lower()
            if st in ("ok", "both", "found"):
                ws.cell(r, len(row)).fill = ok_f
            elif st in ("diff", "missing", "not found", "changed", "no 8949", "data entry", "no document"):
                ws.cell(r, len(row)).fill = bad_f
            elif st.startswith("decision"):
                ws.cell(r, len(row)).fill = PatternFill("solid", fgColor="FFF2CC")
            r += 1
        return r + 1

    ws = wb.active
    ws.title = "Checklist"
    ws["A1"] = f"Review — {client}"
    ws["A1"].font = Font(name=F, size=13, bold=True, color="003366")
    ws["A2"] = f"{cy_file} · saved {datetime.now():%Y-%m-%d %H:%M}"
    ws["A2"].font = norm
    r = table(ws, 4, "Checklist", ["Done", "Item", "Stays (master)", "Note"],
              [["✓" if x.get("done") else "", x["text"], "Y" if x.get("stays") else "N", x.get("note", "")] for x in checklist])
    cm = comments or {}
    if cm.get("client"):
        ws.cell(r, 1, "Client comments").font = Font(name=F, size=11, bold=True, color="003366")
        x = ws.cell(r + 1, 1, cm["client"])
        x.alignment = Alignment(wrap_text=True, vertical="top")
        ws.merge_cells(start_row=r + 1, start_column=1, end_row=r + 1, end_column=4)
        ws.row_dimensions[r + 1].height = max(15, 15 * (cm["client"].count("\n") + 1 + len(cm["client"]) // 140))
        r += 3
    if expected:
        recv = cm.get("recv") or {}
        ok = cm.get("recv_ok") or {}
        r = table(ws, r, "On last year's return — not received", ["Verified", "Form", "Payer / entity", "Comment"],
                  [["✓" if ok.get(f'{e["kind"]}|{e["name"]}') else "", e["kind"], e["name"], recv.get(f'{e["kind"]}|{e["name"]}', "")]
                   for e in expected if not e.get("received")])
        if cm.get("recv_note"):
            ws.cell(r - 1, 1, "Note: " + cm["recv_note"]).font = norm
            r += 1
    for col, w in zip("ABCD", (8, 80, 14, 50)):
        ws.column_dimensions[col].width = w

    if ut:
        ws = wb.create_sheet("UT input check")
        lbl = {"ok": "ok", "red": "data entry", "amber": "decision", "note": "note"}
        r = table(ws, 1, "1 · Document verification — every box on the documents vs the UT entry",
                  ["✓", "Group", "Document", "Box / line", "Per document", "Per preparer", "Per UT", "UT screen", "Comment (summary)",
                   "Review comment", "Status"],
                  [[tick("v|" + x["key"]), x["group"], x["doc"], x["box"], x["per_doc"], x["per_prep"], x["ut"],
                    (x["ut_ref"] + ("; " + x["ut_note"] if x.get("ut_note") else "")), x["note"], mk.get("n|v|" + x["key"], ""),
                    lbl.get(x["status"], x["status"]) + (" — no comment" if x["no_comment"] else "") +
                    ("" if x["entry"] in ("ok", "note") else " (" + x["entry"] + ")")] for x in ut["lines"]])
        r = table(ws, r, "2 · UltraTax entries with no document", ["✓", "UT screen", "Entry", "What", "Amount", "Page", "Review comment", "Status"],
                  [[tick("v|" + x["key"]), x["screen"], x["label"], x.get("what", ""), x["amount"], x["page"], mk.get("n|v|" + x["key"], ""),
                    "no document"] for x in ut["back"]])
        r = table(ws, r, "Payments on the input screens", ["Screen", "Kind", "Line", "Amount", "Page"],
                  [[x["screen"], x["kind"], x["label"], x["amount"], x["page"]] for x in ut["payments"]])
        for col, w in zip("ABCDEFGHIJK", (5, 22, 40, 34, 14, 14, 14, 46, 36, 36, 22)):
            ws.column_dimensions[col].width = w
    if inc:
        ws = wb.create_sheet("Reviewed (✓)")
        rows = []
        for title, summ in (("Income", inc), ("Deductions", ded or {"sections": []})):
            for sct in summ["sections"]:
                for x in sct["rows"]:
                    if mk.get("inc|" + x["id"]):
                        rows.append([title, sct["label"].replace("Total ", ""), x["doc"], x.get("cy"), "line"])
                if mk.get("inctot|" + sct["key"]):
                    rows.append([title, sct["label"], "", sct.get("cy"), "total"])
        rows += [["Other", "", k.split("|", 1)[-1], None, k.split("|", 1)[0]] for k in mk
                 if not k.startswith(("inc|", "inctot|", "v|", "n|")) and "|" in k]
        r = table(ws, 1, "Lines ticked as reviewed (✓) on screen 1", ["Summary", "Category", "Line", "Amount", "Kind"], rows)
        for col, w in zip("ABCDE", (14, 34, 60, 16, 12)):
            ws.column_dimensions[col].width = w
    if checks:
        _return_sheet(wb, checks, table)
    if docs:
        from core.excel import documents_sheet
        documents_sheet(wb, docs)
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _return_sheet(wb, checks, table):
    ws = wb.create_sheet("Return checks")
    r = 1
    r = table(ws, r, "1 · Income by category — documents vs return", ["Category", "Return line", "Documents", "Return", "Difference", "Status"],
              [[x["cat"], x["line"], x["docs"], x["ret"], x["diff"], x["status"]] for x in checks["cats"]])
    r = table(ws, r, "2 · Each amount on the return (program, changed and added)", ["Category", "Document", "Amount", "Source", "Return pages", "Status"],
              [[x["cat"], x["doc"], x["amount"], (x.get("source", "") + (" — " + x["change"] if x.get("change") else "")),
                ", ".join(map(str, x["pages"])), "found" if x["found"] else "not found"] for x in checks["perdoc"]])
    r = table(ws, r, "3 · Schedule D by line / Form 8949 box", ["Schedule D line", "Documents", "Return", "Difference", "Rows", "Status"],
              [[x["line"], x["docs"], x["ret"], x["diff"], "; ".join(x["rows"]), x["status"]] for x in checks["schd"]])
    r = table(ws, r, "4 · Form 8949 boxes", ["8949 page", "Box", "Total (h) / documents", "Note", "Status"],
              [[x.get("page"), x["box"], x.get("total") if x.get("total") is not None else "; ".join(x.get("docs", [])), x.get("note", ""),
                x.get("status", "")] for x in checks["boxes"]])
    r = table(ws, r, "5 · Forms filed — 2024 vs 2025", ["Form", "Short name", "Status"],
              [[x["form"], x["short"], x["status"]] for x in checks["forms"]])
    r = table(ws, r, "Payers / K-1 entities — 2024 vs 2025", ["Form", "Payer / entity", "Status"],
              [[x["kind"], x["name"], x["status"]] for x in checks["payers"]])
    r = table(ws, r, "Other", ["Item", "2024", "2025", "Status"],
              [[x["item"], x.get("py"), x.get("cy"), x["status"]] for x in checks["status"]])
    if checks.get("sche"):
        r = table(ws, r, "Schedule E — each property, each line (box vs return)", ["Property (box)", "On the return", "Line", "Box", "Return", "Status"],
                  [[x["prop"], x["ret_prop"], f'{x["line"]} {x["label"]}', x["box"], x["ret"], x["status"]] for x in checks["sche"]])
    if checks.get("memo"):
        r = table(ws, r, "1095-A / 1098-T (not income)", ["Form", "Document", "What", "Amount", "On the return", "Status"],
                  [[x["form"], x["doc"], x["what"], x["amount"], x["where"] + " · " + x.get("note", ""), x["status"]] for x in checks["memo"]])
    for col, w in zip("ABCDEF", (44, 46, 16, 16, 16, 14)):
        ws.column_dimensions[col].width = w
