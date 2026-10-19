"""masking.py - mask client PII in the documents that go to the AI.

What is masked (from the person tabs on the masking screen):
  last names (first names stay visible), SSNs / ITINs (any format, last 4 digits stay visible), EINs (7-digit suffix),
  addresses, phones, e-mails, account numbers and any other text typed in.
Automatic (only the obvious): SSNs printed as 123-45-6789 (last 4 stay), e-mail addresses (xxx@xxx.xxx) and the ZIP after a state
(+ the street line above it).
Never masked: payer / employer names and dollar amounts (they are what the review needs).

Formats
  PDF   - true redaction, text layer kept (redact_pdf.py engine)
  Word  - .docx text replaced with XXXX (body, tables, headers, footers)
  Excel - .xlsx cell text replaced
  Text  - .txt .csv .eml .htm(l) .md replaced
  Anything else (pictures, .doc, .msg, ...) is NOT masked -> listed with a warning, never saved as "redacted".
"""
from __future__ import annotations

import re
import shutil
from pathlib import Path

from . import redact_pdf as RP

PDF = {".pdf"}
DOCX = {".docx"}
XLSX = {".xlsx", ".xlsm"}
TEXT = {".txt", ".csv", ".eml", ".htm", ".html", ".md", ".json", ".xml"}
MASKABLE = PDF | DOCX | XLSX | TEXT

# automatic: only the obvious - SSN printed as 123-45-6789 and city/state ZIP. Everything else is typed in the tabs.
AUTO_CATS = {"ssn", "address", "email"}   # never "employer" (payer names must stay), never "name" by label
EMAIL = re.compile(r"\b[\w.+\-]+@[\w\-]+(?:\.[\w\-]+)*\.[A-Za-z]{2,}\b")

SSN_DASH = re.compile(r"(?<![\d,.$])\d{3}[-\u2013\u2014]\d{2}[-\u2013\u2014]\d{4}(?![\d,.])")
STATE_ZIP = RP.STATE_ZIP_RE


# always masked, on every document (firm-level identifiers)
from .settings import ALWAYS_MASK                    # from local_settings.json (not in GitHub)

FIELD_KEYS = ["names", "ids", "accounts", "emails", "streets", "cityzip", "other"]
EIN_RE = re.compile(r"^\d{2}[-\u2013\u2014]\d{7}$")


def collect_values(people) -> dict:
    """Masking boxes -> {'name': [...], 'ssn': [...], 'ein': [...], 'other': [...]}

    people = {"names": [...], "ids": [...], "accounts": [...], "emails": [...], "streets": [...], "cityzip": [...],
              "other": [...]} - each list holds the box values and the "add more" text (comma separated).
    (An older list-of-person-tabs format is still accepted.)"""
    out = {"name": [], "ssn": [], "ein": [], "other": []}
    if isinstance(people, dict):
        def vals(k):
            res = []
            for v in people.get(k) or []:
                res += _split(v, sep=",") if "," in str(v) and k not in ("streets", "cityzip") else _split(v, sep="\n")
            return res
        out["name"] = vals("names")
        for v in vals("ids"):
            (out["ein"] if EIN_RE.match(v) else out["ssn"]).append(v)
        for k in ("accounts", "emails", "streets", "other"):
            out["other"] += vals(k)
        for v in ALWAYS_MASK:
            (out["ein"] if EIN_RE.match(v) else out["other"]).append(v)
        for v in vals("cityzip"):                       # "Springfield; 32084" / "Springfield 32084" -> city and ZIP
            for part in re.split(r"[;,]", v):
                part = part.strip()
                m = re.match(r"^(.*?)\s*\b(\d{5}(?:-\d{4})?)$", part)
                if m:
                    out["other"] += [x for x in (m.group(1).strip(" ,"), m.group(2)) if x]
                elif part:
                    out["other"].append(part)
        people = []
    for p in people:
        for ln in _split(p.get("last")):
            out["name"].append(ln)
        out["ssn"] += _split(p.get("ssn"))
        out["ein"] += _split(p.get("ein"))
        for k in ("address", "city", "phone", "email", "accounts", "other"):
            out["other"] += _split(p.get(k), sep=";" if k in ("address", "city") else ",;")
    for k in out:
        seen, vals = set(), []
        for v in out[k]:
            if v.lower() not in seen and len(v) >= 2:
                seen.add(v.lower())
                vals.append(v)
        out[k] = vals
    return out


def _split(v, sep=",;"):
    if not v:
        return []
    return [x.strip() for x in re.split("[" + re.escape(sep) + "\n]", str(v)) if x.strip()]


def _text_patterns(values: dict) -> list:
    """[(pattern, kind)] - kind 'ssn' keeps the last 4 digits, 'zip' keeps the state, 'x' masks it all."""
    pats = []
    for key in ("name", "ssn", "ein", "other"):
        for v in values.get(key, []):
            if key == "ssn":
                d = re.sub(r"\D", "", v)
                if len(d) >= 4:
                    pats.append((re.compile(r"\b" + r"[-\u2013\u2014.\s]?".join(d) + r"\b"), "ssn"))
            elif key == "ein":
                m = re.match(r"^(\d{2})[-–—]?(\d{7})$", v.strip())
                pats.append((re.compile(r"\b" + (m.group(2) if m else re.escape(v)) + r"\b"), "x"))
            else:
                toks = [re.escape(t) for t in v.split()]
                # letters/digits around it end the match, "_" doesn't (file names like US_1040_Smith_2024.pdf)
                pats.append((re.compile(r"(?<![A-Za-z0-9])" + r"\s+".join(toks) + r"(?![A-Za-z0-9])", re.I), "x"))
    pats.sort(key=lambda p: -len(p[0].pattern))
    return pats + [(SSN_DASH, "ssn"), (STATE_ZIP, "zip"), (EMAIL, "x")]


def _keep_last4(t: str) -> str:
    out, seen, total = [], 0, sum(c.isdigit() for c in t)
    for c in t:
        if c.isdigit():
            seen += 1
            out.append("X" if seen <= total - 4 else c)
        else:
            out.append(c)
    return "".join(out)


def mask_text(s: str, pats) -> tuple[str, int]:
    n = 0
    for p, kind in pats:
        if kind == "ssn":
            fn = lambda m: _keep_last4(m.group(0))
        elif kind == "zip":
            fn = lambda m: m.group(1) + " " + "X" * 5
        else:
            fn = lambda m: "X" * min(max(len(m.group(0)), 4), 10)
        s, k = p.subn(fn, s)
        n += k
    return s, n


def _mask_docx(src: Path, dst: Path, pats) -> int:
    import docx
    d = docx.Document(str(src))
    total = 0

    def paras(container):
        for p in container.paragraphs:
            yield p
        for t in getattr(container, "tables", []):
            for row in t.rows:
                for cell in row.cells:
                    yield from paras(cell)

    containers = [d] + [s.header for s in d.sections] + [s.footer for s in d.sections]
    for c in containers:
        for p in paras(c):
            new, n = mask_text(p.text, pats)
            if n:
                total += n
                runs = p.runs
                if runs:
                    runs[0].text = new
                    for r in runs[1:]:
                        r.text = ""
    d.core_properties.author = ""
    d.core_properties.last_modified_by = ""
    d.save(str(dst))
    return total


def _mask_xlsx(src: Path, dst: Path, pats) -> int:
    import openpyxl
    wb = openpyxl.load_workbook(str(src))
    total = 0
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for c in row:
                if isinstance(c.value, str):
                    new, n = mask_text(c.value, pats)
                    if n:
                        c.value, total = new, total + n
    wb.properties.creator = ""
    wb.properties.lastModifiedBy = ""
    wb.save(str(dst))
    return total


def _mask_textfile(src: Path, dst: Path, pats) -> int:
    raw = src.read_bytes()
    for enc in ("utf-8", "cp1252", "latin-1"):
        try:
            s = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    new, n = mask_text(s, pats)
    dst.write_text(new, encoding="utf-8")
    return n


def _trim_pdf(src: Path, pages: list[int], workdir: Path, passwords) -> Path:
    """Copy of a long PDF with only the pages the AI needs (1-based page numbers)."""
    from .pdfutil import open_pdf
    doc, _ = open_pdf(src, passwords)
    keep = [p - 1 for p in pages if 1 <= p <= len(doc)]
    doc.select(keep)
    out = workdir / "_pages" / src.name
    out.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out), garbage=3, deflate=True)
    doc.close()
    return out


def mask_files(files: list[str], values: dict, workdir: Path, passwords=None, pages: dict | None = None) -> list[dict]:
    """Masks each file into workdir (same file name). Returns one result per file.
    pages: {path: [page numbers]} - long PDFs keep only those pages (the summary pages the AI needs)."""
    workdir.mkdir(parents=True, exist_ok=True)
    pats = _text_patterns(values)
    pwd = {"*": list(passwords or [])}
    results = []
    for f in files:
        src = Path(f)
        ext = src.suffix.lower()
        r = dict(file=str(src), name=src.name, ext=ext, status="", masked=None, hits=0, warn="", left=[])
        if not src.is_file():
            r.update(status="missing", warn="File not found.")
            results.append(r)
            continue
        dst = workdir / src.name
        i = 2
        while dst.exists():                         # two documents with the same name
            dst = workdir / f"{src.stem} ({i}){src.suffix}"
            i += 1
        try:
            if ext in PDF:
                keep = (pages or {}).get(str(src))
                srcpdf = _trim_pdf(src, keep, workdir, passwords) if keep else src
                if keep:
                    r["pages"] = keep
                rep = RP.redact_pdf(srcpdf, dst, AUTO_CATS, overrides=values, exclude_set=set(), passwords=pwd)
                r["hits"] = rep["redactions"]
                left = [f"{k[1]} (page {', '.join(str(p) for p, _ in v if p)})" for k, v in rep["verification"].items()
                        if k[0] != "error"]
                r["left"] = left
                notes = []
                if rep["pages_no_text"]:
                    notes.append(f"page(s) {', '.join(map(str, rep['pages_no_text'][:8]))} have no text layer — "
                                 "anything printed there as an image is NOT masked")
                if rep["pages_failed"]:
                    notes.append("masking failed on page(s) " + ", ".join(str(p) for p, _ in rep["pages_failed"]))
                if left:
                    notes.append("still found after masking: " + "; ".join(left))
                if keep:
                    notes.insert(0, f"only pages {', '.join(map(str, keep))} kept (summary pages for the AI)")
                r["warn"] = " · ".join(notes)
                r["status"] = "check" if (left or rep["pages_failed"] or rep["pages_no_text"]) else "masked"
            elif ext in DOCX:
                r["hits"] = _mask_docx(src, dst, pats)
                r["status"] = "masked"
                r["warn"] = "Word: text masked; pictures / text boxes inside the document are not checked."
            elif ext in XLSX:
                r["hits"] = _mask_xlsx(src, dst, pats)
                r["status"] = "masked"
            elif ext in TEXT:
                r["hits"] = _mask_textfile(src, dst, pats)
                r["status"] = "masked"
            else:
                r.update(status="format", warn=f"{ext or 'This'} files can't be masked — NOT included in the redacted set. "
                                               "Convert to PDF (with text) or review by hand.")
                results.append(r)
                continue
            r["masked"] = str(dst)
        except PermissionError as e:
            r.update(status="failed", warn=str(e))
        except Exception as e:  # noqa: BLE001
            r.update(status="failed", warn=f"{type(e).__name__}: {e}")
        results.append(r)
    return results


def save_redacted(results: list[dict], folder: str, name_fn=None) -> list[str]:
    """Copies the masked files as redacted_<name>. name_fn masks the file name too (client names in file names)."""
    out = Path(folder)
    out.mkdir(parents=True, exist_ok=True)
    saved = []
    for r in results:
        if r.get("masked") and r["status"] in ("masked", "check"):
            nm = Path(r["masked"]).name
            dst = _free(out / ("redacted_" + (name_fn(nm) if name_fn else nm)))
            shutil.copy2(r["masked"], dst)
            saved.append(str(dst))
    return saved


def save_copies(files: list[str], folder: str) -> list[str]:
    out = Path(folder)
    out.mkdir(parents=True, exist_ok=True)
    saved = []
    for f in files:
        src = Path(f)
        if src.is_file():
            dst = _free(out / src.name)
            shutil.copy2(src, dst)
            saved.append(str(dst))
    return saved


def _free(p: Path) -> Path:
    i, q = 2, p
    while q.exists():
        q = p.with_name(f"{p.stem} ({i}){p.suffix}")
        i += 1
    return q
