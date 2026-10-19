"""sche_template.py - the standard Schedule E sheet (one row per property) - write a blank / filled one, read it back.

Same layout every time, so the 1040 app (Schedule E box -> Import) and the Schedule E skill agree:
  sheet "Schedule E": row 1 = headers below, one row per property, a TOTAL row at the end (ignored on import)
  sheet "Flags":      what needs the preparer's judgement (written by the skill)
  sheet "How to fill": short instructions
Several Schedule Es: "Schedule E #" = the tab in the 1040 app (6 properties per tab) and "Owner" = T / S / J.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path

# key, header (the header text is what the import matches - keep it)
COLUMNS = [
    ("sch", "Schedule E #"), ("owner", "Owner (T/S/J)"), ("num", "Property #"), ("short", "Short address"),
    ("address", "Full address"), ("type", "Type (1-8)"), ("fair_days", "Fair rental days"), ("personal_days", "Personal use days"),
    ("qjv", "QJV (Y/N)"),
    ("3", "3 Rents received"), ("4", "4 Royalties received"),
    ("5", "5 Advertising"), ("6", "6 Auto and travel"), ("7", "7 Cleaning and maintenance"), ("8", "8 Commissions"),
    ("9", "9 Insurance"), ("10", "10 Legal and other professional fees"), ("11", "11 Management fees"),
    ("12", "12 Mortgage interest paid to banks"), ("13", "13 Other interest"), ("14", "14 Repairs"), ("15", "15 Supplies"),
    ("16", "16 Taxes"), ("17", "17 Utilities"), ("18", "18 Depreciation or depletion"), ("19", "19 Other - total"),
    ("22", "22 Loss allowed after limitation (Form 8582)"),
    ("source", "Source"), ("flags", "Flags / notes"),
]
AMOUNT_KEYS = ["3", "4", "5", "6", "7", "8", "9", "10", "11", "12", "13", "14", "15", "16", "17", "18", "19", "22"]
TYPES = {"1": "Single Family Residence", "2": "Multi-Family Residence", "3": "Vacation/Short-Term Rental", "4": "Commercial",
         "5": "Land", "6": "Royalties", "7": "Self-Rental", "8": "Other"}
HOW_TO = [
    "One row per property. Leave a cell blank when there is no amount - never type 0 for 'unknown'.",
    "Schedule E #: 1 for the first six properties, 2 for the next six, ... (one tab of 6 properties in the 1040 app).",
    "19 Other - total: the total of the other expenses only (no description needed).",
    "Owner: T = taxpayer, S = spouse, J = joint (a separate Schedule E per owner when the return needs it).",
    "Type: 1 Single family, 2 Multi-family, 3 Vacation/short-term, 4 Commercial, 5 Land, 6 Royalties, 7 Self-rental, 8 Other.",
    "Expenses are positive numbers. 18 Depreciation: from the UltraTax asset module - leave blank if unknown.",
    "22: only when the loss is limited by Form 8582 (the allowed loss, positive number).",
    "Source: where the numbers came from (client spreadsheet, 1098 NewRez, tax bill ...). Flags / notes: anything to check.",
    "The TOTAL row is for tying to the client's own totals - the 1040 app ignores it on import.",
]


def _norm(h: str) -> str:
    return re.sub(r"[^a-z0-9]+", " ", str(h or "").lower()).strip()


HEADER_KEY = {_norm(h): k for k, h in COLUMNS}
HEADER_KEY.update({_norm("19 Other"): "19", _norm("19 Other - description"): "19d"})     # older sheets


def write_template(path, rows: list | None = None, flags: list | None = None) -> str:
    """Writes the standard sheet (blank when rows is None). rows: [{key: value}] with the COLUMNS keys."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    ws = wb.active
    ws.title = "Schedule E"
    head, hfill = Font(name="Arial", size=10, bold=True, color="FFFFFF"), PatternFill("solid", fgColor="003366")
    norm, bold = Font(name="Arial", size=10), Font(name="Arial", size=10, bold=True)
    for c, (_, h) in enumerate(COLUMNS, 1):
        x = ws.cell(1, c, h)
        x.font, x.fill = head, hfill
        x.alignment = Alignment(wrap_text=True, vertical="top")
    rows = rows or []
    for r, row in enumerate(rows, 2):
        for c, (k, _) in enumerate(COLUMNS, 1):
            v = row.get(k)
            x = ws.cell(r, c, v if v not in ("",) else None)
            x.font = norm
            if k in AMOUNT_KEYS:
                x.number_format = '#,##0.00;-#,##0.00;""'
    last = max(2, len(rows) + 1)
    tot = last + 1
    ws.cell(tot, 4, "TOTAL").font = bold
    for c, (k, _) in enumerate(COLUMNS, 1):
        if k in AMOUNT_KEYS:
            col = get_column_letter(c)
            x = ws.cell(tot, c, f"=SUM({col}2:{col}{last})")
            x.font, x.number_format = bold, '#,##0.00;-#,##0.00;""'
    widths = {"short": 26, "address": 36, "source": 30, "flags": 50}
    for c, (k, _) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(c)].width = widths.get(k, 13)
    ws.row_dimensions[1].height = 42
    ws.freeze_panes = "E2"
    fl = wb.create_sheet("Flags")
    for c, h in enumerate(["Property", "Line", "Flag"], 1):
        x = fl.cell(1, c, h)
        x.font, x.fill = head, hfill
    for r, f in enumerate(flags or [], 2):
        fl.cell(r, 1, f.get("property", "")).font = norm
        fl.cell(r, 2, f.get("line", "")).font = norm
        fl.cell(r, 3, f.get("flag", "")).font = norm
    for col, w in zip("ABC", (30, 14, 100)):
        fl.column_dimensions[col].width = w
    ht = wb.create_sheet("How to fill")
    for r, t in enumerate(HOW_TO, 1):
        ht.cell(r, 1, t).font = norm
    ht.column_dimensions["A"].width = 120
    wb.save(str(path))
    return str(path)


def _num(v):
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("$", "").replace(",", "")
    if s in ("", "-", "—"):
        return None
    neg = s.startswith("(") and s.endswith(")")
    try:
        x = float(s.strip("()"))
    except ValueError:
        return None
    return -x if neg else x


def _rows_from(path) -> list[list]:
    p = Path(path)
    if p.suffix.lower() == ".csv":
        with open(p, newline="", encoding="utf-8-sig") as f:
            return [r for r in csv.reader(f)]
    from openpyxl import load_workbook
    wb = load_workbook(str(p), data_only=True, read_only=True)
    ws = wb["Schedule E"] if "Schedule E" in wb.sheetnames else wb.worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


def read_template(path) -> dict:
    """-> {properties: [{name, sch, owner, num, info: {...}, cy: {line: amount}, notes}], errors: [...]}"""
    rows = _rows_from(path)
    hi = next((i for i, r in enumerate(rows[:10]) if sum(1 for c in r if _norm(c) in HEADER_KEY) >= 5), None)
    if hi is None:
        return dict(properties=[], errors=["This isn't the Schedule E template (headers not found)."])
    keys = [HEADER_KEY.get(_norm(c)) for c in rows[hi]]
    props, errors = [], []
    for n, r in enumerate(rows[hi + 1:], hi + 2):
        rec = {k: v for k, v in zip(keys, r) if k}
        label = " ".join(str(rec.get(k) or "") for k in ("short", "address")).strip()
        if not label and not any(_num(rec.get(k)) for k in AMOUNT_KEYS):
            continue
        if re.match(r"^\s*total\b", label, re.I) or re.match(r"^\s*total\b", str(rec.get("num") or ""), re.I):
            continue
        cy = {}
        for k in AMOUNT_KEYS:
            v = _num(rec.get(k))
            if v is not None:
                cy[k] = abs(v) if k == "22" else v
        t = str(rec.get("type") or "").strip()
        t = TYPES.get(t[:1], t) if t[:1].isdigit() else t
        props.append(dict(
            name=str(rec.get("short") or rec.get("address") or f"Property {len(props) + 1}").strip(),
            sch=int(_num(rec.get("sch")) or 0) or None, owner=str(rec.get("owner") or "").strip().upper()[:1] or "T",
            num=int(_num(rec.get("num")) or 0) or None, cy=cy, notes=str(rec.get("flags") or "").strip(),
            info=dict(type=t, fair_days=_num(rec.get("fair_days")), personal_days=_num(rec.get("personal_days")),
                      qjv=str(rec.get("qjv") or "").strip().upper()[:1] == "Y", other_desc=str(rec.get("19d") or "").strip(),
                      full_address=str(rec.get("address") or "").strip(), source=str(rec.get("source") or "").strip(),
                      owner=str(rec.get("owner") or "").strip().upper()[:1] or "T")))
        if not props[-1]["cy"]:
            errors.append(f"Row {n} ({props[-1]['name']}): no amounts")
    # tab numbers: as given, else 6 per tab in order
    for i, p in enumerate(props):
        p["sch"] = p["sch"] or i // 6 + 1
    props.sort(key=lambda p: ({"T": 0, "S": 1, "J": 2}.get(p["owner"], 3), p["sch"], p["num"] or 0))
    return dict(properties=props, errors=errors)
